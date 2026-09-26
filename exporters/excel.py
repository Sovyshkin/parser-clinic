from __future__ import annotations

import asyncio
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from database.db import ClinicRepository
from database.models import Clinic


HEADERS = (
    "Domain",
    "Website",
    "Clinic Name",
    "Phones",
    "Emails",
    "Telegram",
    "WhatsApp",
    "VK",
    "Address",
    "Parsed At",
)


class ExcelExporter:
    def __init__(self, repository: ClinicRepository, path: Path) -> None:
        self.repository = repository
        self.path = path
        self._lock = asyncio.Lock()

    async def export_all_to_excel(self) -> Path:
        async with self._lock:
            clinics = await self.repository.list_parsed()
            await asyncio.to_thread(self._write, clinics)
        return self.path

    def _write(self, clinics: list[Clinic]) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Clinics"
        sheet.sheet_view.showGridLines = False
        sheet.append(HEADERS)

        for clinic in clinics:
            parsed_at: datetime | str | None = clinic.parsed_at
            if parsed_at:
                try:
                    parsed_at = datetime.fromisoformat(parsed_at)
                    if parsed_at.tzinfo is not None:
                        parsed_at = parsed_at.astimezone(timezone.utc).replace(tzinfo=None)
                except ValueError:
                    pass
            sheet.append(
                (
                    clinic.domain,
                    clinic.website,
                    clinic.clinic_name,
                    "; ".join(clinic.phones),
                    "; ".join(clinic.emails),
                    "; ".join(clinic.telegram),
                    "; ".join(clinic.whatsapp),
                    "; ".join(clinic.vk),
                    clinic.address,
                    parsed_at,
                )
            )

        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top")
        for cell in sheet["J"][1:]:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"

        widths = (24, 38, 32, 28, 30, 30, 30, 30, 48, 22)
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:J{max(sheet.max_row, 1)}"
        if sheet.max_row >= 2:
            table = Table(displayName="ClinicsTable", ref=f"A1:J{sheet.max_row}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)

        temp_path = self.path.with_suffix(".tmp.xlsx")
        workbook.save(temp_path)
        temp_path.replace(self.path)


async def export_all_to_excel(repository: ClinicRepository, path: Path) -> Path:
    return await ExcelExporter(repository, path).export_all_to_excel()
