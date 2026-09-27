import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from database.db import ClinicRepository
from database.models import Clinic, ClinicStatus
from exporters.excel import ExcelExporter


class PersistenceAndExportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        self.tempdir = tempfile.TemporaryDirectory()
        root = Path(self.tempdir.name)
        self.db_path = root / "clinics.db"
        self.xlsx_path = root / "clinics.xlsx"
        self.repository = ClinicRepository(self.db_path)
        await self.repository.initialize()

    async def asyncTearDown(self) -> None:
        self.tempdir.cleanup()

    async def test_persistent_domain_deduplication_and_export(self) -> None:
        self.assertTrue(await self.repository.reserve_new("clinic.ru", "https://clinic.ru/"))
        self.assertFalse(await self.repository.reserve_new("clinic.ru", "https://clinic.ru/other"))
        self.assertTrue(await self.repository.reserve_new("pending.ru", "https://pending.ru/"))
        await self.repository.mark_parsed(
            Clinic(
                domain="clinic.ru",
                website="https://clinic.ru/",
                clinic_name="Клиника",
                phones=("+79123456789",),
                emails=("info@clinic.ru",),
                parsed_at="2026-09-26T10:00:00+00:00",
                status=ClinicStatus.PARSED,
            )
        )

        reopened = ClinicRepository(self.db_path)
        await reopened.initialize()
        self.assertEqual(await reopened.get_status("clinic.ru"), ClinicStatus.PARSED)
        self.assertEqual(await reopened.get_status("pending.ru"), ClinicStatus.FAILED)

        path = await ExcelExporter(reopened, self.xlsx_path).export_all_to_excel()
        workbook = load_workbook(path)
        sheet = workbook["Clinics"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["A2"].value, "clinic.ru")
        self.assertEqual(sheet["D2"].value, "+79123456789")
        workbook.close()

    async def test_latest_run_export_excludes_previous_runs(self) -> None:
        first_run_id = await self.repository.start_run("first")
        await self.repository.reserve_new("first.ru", "https://first.ru/")
        await self.repository.mark_parsed(
            Clinic(
                domain="first.ru",
                website="https://first.ru/",
                clinic_name="Первая клиника",
                status=ClinicStatus.PARSED,
            ),
            first_run_id,
        )

        second_run_id = await self.repository.start_run("second")
        await self.repository.reserve_new("second.ru", "https://second.ru/")
        await self.repository.mark_parsed(
            Clinic(
                domain="second.ru",
                website="https://second.ru/",
                clinic_name="Вторая клиника",
                status=ClinicStatus.PARSED,
            ),
            second_run_id,
        )

        path = await ExcelExporter(
            self.repository,
            self.xlsx_path,
        ).export_latest_run_to_excel()
        workbook = load_workbook(path)
        sheet = workbook["Clinics"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["A2"].value, "second.ru")
        workbook.close()


if __name__ == "__main__":
    unittest.main()
